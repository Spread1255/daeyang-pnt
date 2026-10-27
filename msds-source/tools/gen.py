# MSDS 일괄 생성기 (IV064 KOSHA 양식 틀 + 원료 사전 + KOSHA 자료)
import openpyxl, datetime, sys, re, glob, json, html, copy
sys.path.insert(0,'.'); sys.path.insert(0,'/tmp/claude-0/ral')
from xlhelp import *
from ex import txt, sections

def nkey(n): return re.sub(r'[\s#]','',re.sub(r'\(IN\)$','',str(n).strip().upper()))
ALIAS={  # 원료명 → CAS
 'HC7803':'27923-68-8','HC6011':'114267-10-6','HC5803':'114267-10-6','HC6805':'114267-10-6','PC2802':'40471-09-8',
 'YD013K':'25036-25-3','E-12':'25036-25-3','RICHGOLD(CU)':'7440-50-8','RICHGOLD(ZN)':'7440-66-6','PALEGOLD(CU)':'7440-50-8','PALEGOLD(ZN)':'7440-66-6','CC2682':'40471-09-8','QM1260':'57843-53-5','S7000':'1317-65-3','CTCOPPER':'7440-50-8','KER3002':'25036-25-3','KER3003':'25036-25-3','XG665':'37953-21-2',
 'ADP964':'25214-38-4','XG603-1A':'54553-90-1','PCA-501':'6334-25-4','XL552':'6334-25-4','XL-552':'6334-25-4',
 'LA-100':'9003-49-0','L-88':'9003-49-0','BENZOIN':'119-53-9','JS-250':'9002-88-4','CERA-969':'9002-88-4',
 'VISCOGEL':'68953-58-2','PANGELB20':'63800-37-3','AHA-2139':'9010-77-9','CT-0230':'60621-79-6','SONOX626':'26741-53-7',
 'LS-8K':'21645-51-2','S3000':'1317-65-3','SN7000':'1317-65-3','KCB-02':'13462-86-7','BSP-L':'7727-43-7',
 'FR-767':'13463-67-7','HIBLACK50L':'1333-86-4','BLACKS-160':'1333-86-4','MONARCH880':'1333-86-4',
 'RED115':'1309-37-1','RED170':'2786-76-7','YELLOW4000':'5567-15-7','YELLOW420':'51274-00-1','YELLOWL500':'5468-75-7',
 'YELLOWP210':'8007-18-9','YELLOWP261':'68186-90-3','ORANGE920P':'15793-73-4','BLUE515303':'147-14-8','VIOLET23R':'6358-30-1',
 'MARINEBLUE2504':'57455-37-5','MARINEBLUEH55':'57455-37-5','REFLEXALIV':'7429-90-5','PC100':'7429-90-5','PCR214':'7429-90-5'}
DFA={'CERA-994','HDKV15'}

TPL='/tmp/claude-0/ral/iv064.xlsx'
_t=openpyxl.load_workbook(TPL)['Sheet1']
TROW={_t.cell(r,4).value:[_t.cell(r,1).value,_t.cell(r,3).value] for r in range(38,50)}  # IV064 3번 표 (CAS→[물질명, 이명])
NAME3={**TROW,
 '40471-09-8':['폴리에스테르 수지','1,4-벤젠다이카복실산, 2,2-다이메틸-1,3-프로판다이올, 1,2-에탄다이올 및 헥산다이오산 함유 1,3-벤젠다이카복실산 중합체'],
 '6334-25-4':['폴리에스테르 경화제',"N,N,N',N'-Tetrakis(2-hydroxyethyl)adipamide"],
 '9010-77-9':['아크릴산-에틸렌 공중합체(ACRYLIC ACID-ETHYLENE COPOLYMER)','2-프로페노 산, 중합물 ,함유 에텐(2-PROPENOIC ACID, POLYMER WITH ETHENE);'],
 '21645-51-2':['수산화알루미늄',None],
 '26741-53-7':['3,9-비스(2,4-비스(1,1-다이메틸에틸)펜옥시)-2,4,8,10-테트라옥사-3,9-디포스파스피로[5.5]언데칸','SONOX626'],
 '147-14-8':['구리 프탈로시아닌','구리 프탈로시아닌 청색(COPPER PHTHALOCYANINE BLUE);'],
 '6358-30-1':['카브아졸 다이옥사진 바이올렛(CARBAZOLE DIOXAZINE VIOLET)',"디인돌로(3,2-B:3'',2''-M)트리페노디옥사진, 8,18-디클로로-5,15-디에틸-"],
 '1333-86-4':['카본블랙',None],
 '25214-38-4':['폴리에스테르 수지','1,3-Benzenedicarboxylic acid, polymer with 1,4-benzenedicarboxylic acid and 2,2-dimethyl-1,3-propanediol'],
 '37953-21-2':['아크릴 수지','Styrene-butyl acrylate-glycidyl methacrylate-methyl methacrylate copolymer'],
 '60621-79-6':['아크릴 공중합체','Butyl acrylate-methyl methacrylate-styrene-glycidyl methacrylate-α-methyl styrene copolymer'],
 # 경화제: 물질명은 '에폭시/폴리에스테르 경화제', 화학명은 이명 칸 (기존 BLACK 무광 MSDS 기준)
 '54553-90-1':['에폭시 경화제','benzene-1,2,4,5-tetracarboxylic acid, compound with 4,5-dihydro-2-phenyl-1H-imidazole (1:1)'],
 '57843-53-5':['폴리에스테르 경화제','N,N,N,N-테트라키스(2-하이드록시프로필)아디프아미이드'],
 '7440-50-8':['구리',None],'7440-66-6':['아연','Zinc'],'7429-90-5':['알루미늄','Aluminum'],
}
TLABEL={'에폭시 수지':'25036-25-3','에폭시수지':'25036-25-3','벤조인':'119-53-9','폴리에틸렌(POLYETHYLENE)':'9002-88-4','Barite (Ba(SO4))':'13462-86-7',
 '이산화티타늄':'13463-67-7','C.I. 색소 황색 83(C.I. PIGMENT YELLOW 83)':'5567-15-7','Limestone':'1317-65-3','산화철':'1309-37-1'}

# ---------- KOSHA 자료 ----------
KC={}
def kosha(cas):
    if cas in KC: return KC[cas]
    fs=glob.glob(f'/tmp/claude-0/ral/d_{cas}_*.html')
    KC[cas]=sections(txt(fs[0]))[1] if fs else None
    return KC[cas]
def clean(l): return html.unescape(l.replace('\xa0',' ')).strip()
def vals_between(lines,start,stops):
    out=[];on=False
    for l in lines:
        l=clean(l)
        if on and (l in stops or re.match(r'^[가-하]\.\s',l) or l=='TOP▲'): break
        if on: out.append(l)
        if l==start: on=True
    return out
def summarize(v):
    v=[x for x in v if x and x not in ('자료없음','해당없음','(자료없음)','-') and not x.startswith('※') and x!='-->']
    if not v: return None
    t=v[0]
    if len(v)>1 and v[1].startswith('(') and len(t)+len(v[1])<180: t+=' '+v[1]
    t=re.sub(r'\s+',' ',t).replace('실험종 : ','')
    if len(t)>200:
        cut=t[:200]; p=max(cut.rfind('. '),cut.rfind('함 '),cut.rfind('음 '))
        t=cut[:p+1] if p>60 else cut.rsplit(' ',1)[0]
    return t
L11=['경구','경피','흡입','피부부식성 또는 자극성','심한 눈손상 또는 자극성','호흡기과민성','피부과민성','발암성','산업안전보건법','고용노동부고시','IARC','OSHA','ACGIH','NTP','EU CLP','환경부','NITE','생식세포변이원성','생식독성','특정 표적장기 독성 (1회 노출)','특정 표적장기 독성 (반복 노출)','흡인유해성','기타 유해성 영향']
L12=['어류','갑각류','조류','잔류성','분해성','농축성','생분해성']
def k11(cas,label):
    s=kosha(cas); 
    if not s: return None
    return summarize(vals_between(s.get(11,[]),label,set(L11)))
def k12(cas,label):
    s=kosha(cas)
    if not s: return None
    return summarize(vals_between(s.get(12,[]),label,set(L12)))
def k8(cas):
    s=kosha(cas)
    if not s: return [],[]
    l=[clean(x) for x in s.get(8,[])]
    dom=vals_between(l,'국내규정',{'ACGIH 규정'}); acg=vals_between(l,'ACGIH 규정',{'생물학적 노출기준'})
    f=lambda a:[re.sub(r'(mg/m3|㎎/㎥)(?=[가-힣(])',r'\1 ',re.sub(r'^TWA\s*:\s*','TWA - ',x)) for x in a if x.startswith(('TWA','STEL','C '))]
    return f(dom),f(acg)
def k15(cas):
    s=kosha(cas)
    if not s: return [],False
    l=[clean(x) for x in s.get(15,[])]
    i=next((k for k,x in enumerate(l) if x.startswith('가.')),None); j=next((k for k,x in enumerate(l) if x.startswith('나.')),None)
    osh=[x for x in l[i+1:j] if x not in ('해당없음','자료없음')] if i is not None and j else []
    w=next((k for k,x in enumerate(l) if x.startswith('마.') and '폐기물' in x),None)
    waste= w is not None and l[w+1]=='지정폐기물'
    return osh,waste
def kname(cas):
    s=kosha(cas)
    if not s: return None
    s3=[clean(x) for x in s[3]]
    nm=s3[2] if len(s3)>2 else None
    al=s3[4] if len(s3)>4 and s3[3]=='이명(관용명)' and s3[4] not in ('CAS 번호','자료없음') else None
    return [nm,al]

COLOR={'BE':'베이지색','BK':'흑색','BL':'청색','BR':'갈색','GY':'회색','IV':'아이보리색','RD':'적색','WH':'백색','YL':'황색','GD':'금색','GN':'녹색','SI':'은색'}

def build(p,no,out):
    """p: dict(code,name,items=[[no,raw,val],...]) → (out xlsx, report)"""
    rep={'code':p['code'],'excluded':[],'unknown':[],'rows':[]}
    comp=[]  # [cas, raw names, value]
    for no_,raw,val in p['items']:
        if no_=='DFA' or nkey(raw) in DFA or val is None: continue
        if nkey(raw) in ('CERA-994','HDKV15'): continue
        cas=ALIAS.get(nkey(raw))
        if not cas: rep['unknown'].append(raw); continue
        if not KEEP(val): rep['excluded'].append((raw,val)); continue
        hit=next((c for c in comp if c[0]==cas),None)
        if hit: hit[1].append(raw); hit[2]=round(hit[2]+val,p.get('nd',6))
        else: comp.append([cas,[raw],float(val)])
    if rep['unknown']: return None,rep
    wb=openpyxl.load_workbook(TPL)
    if '호환성 보고서' in wb.sheetnames: del wb['호환성 보고서']
    ws=wb['Sheet1']
    n=len(comp)
    if n>12: insert_before(ws,50,[(None,None)]*(n-12),tpl=49)
    elif n<12:
        for r in range(49,37+n,-1): delete_row(ws,r)
    present=set()
    for i,(cas,raws,v) in enumerate(comp):
        nm=NAME3.get(cas) or kname(cas) or [raws[0],None]
        r=38+i
        ws.cell(r,1).value=nm[0]; ws.cell(r,3).value=nm[1]; ws.cell(r,4).value=cas; ws.cell(r,5).value=p.get('rng',{}).get(cas,rng_item(cas,v))
        L=max(len(nm[0] or ''),len(nm[1] or '')); ws.row_dimensions[r].height=56.25 if L>90 else 45 if L>60 else (33.75 if L>34 else 24.75)
        present.add(cas); rep['rows'].append((', '.join(raws),cas,v,p.get('rng',{}).get(cas,rng_item(cas,v))))
    short={}
    for cas,raws,v in comp:
        nm=(NAME3.get(cas) or kname(cas) or [raws[0]])[0]
        tl=[l for l,c in TLABEL.items() if c==cas and l!='에폭시수지']
        short[cas]=tl[0] if tl else (nm if len(nm)<=30 else raws[0])
    # 1·9·16번
    name=p['name'].strip(); code=p['code'].strip()
    ws['B3']=name; ws['B6']=f'{name} ({code})'
    ws['B8']='에폭시 폴리에스테르계 분체도료 일반용' if name.startswith('내장') else '폴리에스테르계 분체도료 일반용'
    ws['A4'].value=ws['A4'].value.replace('AA14944-0000000102',no)
    r=row_of(ws,lambda t:t=='색상',*sec_range(ws,9)); ws.cell(r,2).value=COLOR.get(code.split('-')[1][:2],ws.cell(r,2).value)
    a16,b16=sec_range(ws,16); r=row_of(ws,lambda t:t.startswith('나. 최초작성일'),a16,b16)
    ws.cell(r,3).value=datetime.datetime(2024,5,9)
    ws.cell(row_of(ws,lambda t:t=='개정횟수',r,ws.max_row+1),3).value='1 회'
    ws.cell(row_of(ws,lambda t:t=='최종개정일자',r,ws.max_row+1),3).value=datetime.datetime(2026,1,6)
    # 행 서식 견본 (삭제 전에 저장)
    SNAP={}
    for secn in (8,11,12,15):
        a,b=sec_range(ws,secn)
        k=next(k for k in range(a,b) if str(ws.cell(k,1).value or '').startswith(' '*15) and ws.cell(k,2).value)
        SNAP[secn]=snap_row(ws,k)
    # 없는 성분의 틀 행 삭제 (이어지는 값 행 포함)
    for secn in (15,12,11,8):
        a,b=sec_range(ws,secn); dels=[]
        for k in range(a,b):
            cas=TLABEL.get(norm(ws.cell(k,1).value))
            if cas and cas not in present:
                dels.append(k); j=k+1
                while j<b and not norm(ws.cell(j,1).value) and ws.cell(j,2).value: dels.append(j); j+=1
        for k in sorted(set(dels),reverse=True): delete_row(ws,k)
    # 틀에 없는 성분 행 추가 (KOSHA)
    def has_child(lo,hi,cas):
        lab={l for l,c in TLABEL.items() if c==cas}
        return any(norm(ws.cell(k,1).value) in lab for k in range(lo,hi))
    def block(label,lo,hi,stop):
        s=row_of(ws,lambda t:t==label,lo,hi); e=s+1
        while e<hi and not stop(norm(ws.cell(e,1).value)): e+=1
        return s,e
    OPS=[]
    a8,b8=sec_range(ws,8)
    for lab,nxt,idx in (('국내규정','ACGIH 규정',0),('ACGIH 규정','나. 적절한 공학적 관리',1)):
        s,e=block(lab,a8,b8,lambda t,nxt=nxt:t==nxt)
        items=[]
        for cas,raws,v in comp:
            if has_child(s+1,e,cas): continue
            for x in k8(cas)[idx]: items.append((short[cas],x))
        if items: OPS.append((e,items,s))
    a11,b11=sec_range(ws,11)
    lab11={'- 경구':'경구','- 경피':'경피','- 흡입':'흡입','- 피부부식성 또는 자극성':'피부부식성 또는 자극성','- 심한 눈손상 또는 자극성':'심한 눈손상 또는 자극성',
           '- 피부과민성':'피부과민성','- 고용노동부고시':'고용노동부고시','- IARC':'IARC','- OSHA':'OSHA','- ACGIH':'ACGIH','- NTP':'NTP','- EU CLP':'EU CLP',
           '- 생식세포변이원성':'생식세포변이원성','- 생식독성':'생식독성','- 특정 표적장기 독성 (1회 노출)':'특정 표적장기 독성 (1회 노출)','- 특정 표적장기 독성 (반복 노출)':'특정 표적장기 독성 (반복 노출)'}
    isl=lambda t: t.startswith(('-','*','나.','다.','라.','마.'))
    for lab,kl in lab11.items():
        s,e=block(lab,a11,b11,isl)
        items=[]
        for cas,raws,v in comp:
            if has_child(s+1,e,cas): continue
            t=k11(cas,kl)
            if t: items.append((short[cas],t))
        if items:
            # 들여쓰기가 다른 마지막 행(예: 황색 83) 앞에 끼움
            ins=e
            while ins-1>s and len(str(ws.cell(ins-1,1).value or ''))-len(str(ws.cell(ins-1,1).value or '').lstrip())<15 and norm(ws.cell(ins-1,1).value): ins-=1
            OPS.append((ins if ins>s else e,items,s))
    a12,b12=sec_range(ws,12)
    for lab,kl in (('- 어류','어류'),('- 갑각류','갑각류'),('- 조류','조류'),('- 잔류성','잔류성')):
        s,e=block(lab,a12,b12,isl)
        items=[(short[cas],k12(cas,kl)) for cas,raws,v in comp if not has_child(s+1,e,cas) and k12(cas,kl)]
        if items: OPS.append((e,items,s))
    a15,b15=sec_range(ws,15)
    s,e=block('가. 산업안전보건법에 의한 규제',a15,b15,lambda t:t.startswith('나.'))
    items=[]
    for cas,raws,v in comp:
        if has_child(s+1,e,cas): continue
        osh,_=k15(cas)
        for i,x in enumerate(osh): items.append((short[cas] if i==0 else None,x))
    if items: OPS.append((e,items,s))
    s,e=block('라. 위험물안전관리법에 의한 규제',a15,b15,lambda t:t.startswith('바.'))
    items=[(short[cas],'지정폐기물') for cas,raws,v in comp if not has_child(s+1,e,cas) and k15(cas)[1]]
    if items: OPS.append((e,items,s))
    def secof(r): return next(n for n in (8,11,12,15) if sec_range(ws,n)[0]<=r<sec_range(ws,n)[1])
    for row,items,lab in sorted(OPS,key=lambda o:o[0],reverse=True):
        if norm(ws.cell(lab,2).value)=='자료없음': ws.cell(lab,2).value=None
        insert_before(ws,row,items,tpl=SNAP[secof(lab)])
    # 비어 버린 항목은 '자료없음'
    for secn in (8,11,12,15):
        a,b=sec_range(ws,secn)
        for k in range(a+1,b):
            t=norm(ws.cell(k,1).value)
            if (t.startswith('-') or t in ('국내규정','ACGIH 규정')) and not ws.cell(k,2).value:
                nx=norm(ws.cell(k+1,1).value) if k+1<b else '-'
                child=(k+1<b) and ws.cell(k+1,2).value and not (nx.startswith(('-','*')) or nx in ('ACGIH 규정',) or re.match(r'^[가-하]\.',nx))
                if not child: ws.cell(k,2).value='자료없음'
    apply_sec2(ws,p.get('carc','2'))
    wb.save(out); return out,rep
