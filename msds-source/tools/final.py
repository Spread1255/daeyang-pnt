import openpyxl, sys
sys.path.insert(0,'.')
from xlhelp import apply_sec2
wb=openpyxl.load_workbook('DY#1-BL141_MSDS_초안.xlsx'); ws=wb['Sheet1']
def a(i): x=ws._images[i].anchor; return (x._from.colOff,x._from.rowOff,x.to.colOff,x.to.rowOff)
ws['A4'].value=ws['A4'].value.replace('0000000102','0000000103')
ws.row_dimensions[39].height=45
import datetime
from xlhelp import row_of, sec_range
a16,b16=sec_range(ws,16); r=row_of(ws,lambda t:t.startswith('나. 최초작성일'),a16,b16)
ws.cell(r,3).value=datetime.datetime(2024,5,9)
ws.cell(row_of(ws,lambda t:t=='최종개정일자',r,ws.max_row+1),3).value=datetime.datetime(2026,1,6)
apply_sec2(ws)
wb.save('DY#1-BL141_MSDS.xlsx')
