import openpyxl, datetime, sys
sys.path.insert(0,'.')
from xlhelp import *
wb=openpyxl.load_workbook('/tmp/claude-0/ral/iv064.xlsx')
if '호환성 보고서' in wb.sheetnames: del wb['호환성 보고서']
ws=wb['Sheet1']
NAME='외장형 EAGON WHITE 반광'; CODE='DY#2-WH029'; NO='AA14944-0000000134'
T={r:[ws.cell(r,c).value for c in (1,3,4)] for r in range(38,50)}
pba,ben,tio,lime,pb29,y83,fe=T[41],T[42],T[46],T[45],T[47],T[48],T[49]
SXN='3,9-비스(2,4-비스(1,1-다이메틸에틸)펜옥시)-2,4,8,10-테트라옥사-3,9-디포스파스피로[5.5]언데칸'
NEW=[ # 26.10.1 작업지시서 NO 순서 그대로 (DFA 제외)
 ('폴리에스테르 수지','1,4-벤젠다이카복실산, 2,2-다이메틸-1,3-프로판다이올, 1,2-에탄다이올 및 헥산다이오산 함유 1,3-벤젠다이카복실산 중합체','40471-09-8',57),   # PC2802
 ("N,N,N',N'-Tetrakis(2-hydroxyethyl)adipamide",None,'6334-25-4',3),                                   # PCA-501
 (pba[0],pba[1],pba[2],0.85),                                                                           # LA-100
 (ben[0],ben[1],ben[2],0.3),                                                                            # BENZOIN
 ('아크릴산-에틸렌 공중합체(ACRYLIC ACID-ETHYLENE COPOLYMER)','2-프로페노 산, 중합물 ,함유 에텐(2-PROPENOIC ACID, POLYMER WITH ETHENE);','9010-77-9',3.4),  # AHA-2139
 (SXN,'SONOX626','26741-53-7',0.2),                                                                     # SONOX626
 ('수산화알루미늄',None,'21645-51-2',7.56),                                                              # LS-8K
 (lime[0],lime[1],lime[2],4.8),                                                                         # S#3000
 (tio[0],tio[1],tio[2],22.5),                                                                           # FR-767
 ('카본블랙',None,'1333-86-4',0.0262),                                                                   # HI BLACK 50L
 (y83[0],y83[1],y83[2],0.1392),                                                                         # YELLOW 4000
 (fe[0],fe[1],fe[2],0.05),                                                                              # RED 115
]
assert abs(sum(n[3] for n in NEW)-99.8254)<1e-9 and len(NEW)==12
# 사용자 지정 함유량 범위 (CAS별)
RANGE={'40471-09-8':'55 ∽ 59','6334-25-4':'2 ∽ 4','9003-49-0':'0.7 ∽ 0.9','119-53-9':'0.2 ∽ 0.4','9010-77-9':'3 ∽ 4',
 '26741-53-7':'0.1 ∽ 0.3','21645-51-2':'7 ∽ 9','1317-65-3':'4 ∽ 6','13463-67-7':'21 ∽ 23','1333-86-4':'0.02 ∽ 0.03',
 '5567-15-7':'0.12 ∽ 0.14','1309-37-1':'0.04 ∽ 0.06'}
for i,(a,c,d,p) in enumerate(NEW):
    r=38+i; ws.cell(r,1).value=a; ws.cell(r,3).value=c; ws.cell(r,4).value=d; ws.cell(r,5).value=RANGE.get(d) or rng_item(d,p)
    L=max(len(a),len(c or ''))
    ws.row_dimensions[r].height=45 if L>60 else (33.75 if L>34 else 24.75)
# 1번·표지
ws['B3']=NAME; ws['B6']=f'{NAME} ({CODE})'
ws['B8']='폴리에스테르계 분체도료 일반용'
ws['A4'].value=ws['A4'].value.replace('AA14944-0000000102',NO)
# 9번 색상
r=row_of(ws,lambda t:t=='색상',*sec_range(ws,9)); ws.cell(r,2).value='백색'
# 16번
r=row_of(ws,lambda t:t.startswith('나. 최초작성일'),*sec_range(ws,16)); today=datetime.datetime(2026,10,7)
ws.cell(r,3).value=datetime.datetime(2024,5,9)
ws.cell(row_of(ws,lambda t:t=='개정횟수',r,ws.max_row+1),3).value='1 회'
ws.cell(row_of(ws,lambda t:t=='최종개정일자',r,ws.max_row+1),3).value=datetime.datetime(2026,1,6)

# ---- 행 삭제/삽입 (아래에서 위로) ----
OPS=[]
a8,b8=sec_range(ws,8); a11,b11=sec_range(ws,11); a12,b12=sec_range(ws,12); a15,b15=sec_range(ws,15)
GONE={'에폭시 수지','에폭시수지','폴리에틸렌(POLYETHYLENE)','Barite (Ba(SO4))'}
for lo,hi in ((a11,b11),(a12,b12),(a15,b15)):
    for k in range(lo,hi):
        if norm(ws.cell(k,1).value) in GONE: OPS.append((k,'del',None))
def before(label,lo,hi,items): OPS.append((row_of(ws,lambda t:t==label,lo,hi),'ins',items))
CB='카본블랙'; AL='수산화알루미늄'; FE='산화철'; SX='SONOX626'; AEA='아크릴산-에틸렌 공중합체'; HAA="N,N,N',N'-Tetrakis(2-hydroxyethyl)adipamide"
before('ACGIH 규정',a8,b8,[(AL,'TWA -  10mg/m3   알루미늄금속분진'),(FE,'TWA -  5mg/m3'),(CB,'TWA -  3.5mg/m3   발암성 2, 흡입성')])
before('나. 적절한 공학적 관리',a8,b8,[(AL,'TWA - 1  ㎎/㎥'),(CB,'TWA - 3  ㎎/㎥')])
OPS.append((row_of(ws,lambda t:t.startswith('C.I. 색소 황색 83'),a11,b11),'ins',[(AL,'LD50 > 2000 ㎎/㎏ Rat (OECD TG 423, GLP)'),(SX,'LD50 > 5500 ㎎/㎏ Rat'),(FE,'LD50 > 10000 ㎎/㎏ Rat'),(CB,'LD50 > 8000 ㎎/㎏ Rat (OECD TG 401)')]))
before('- 흡입',a11,b11,[(SX,'LD50 ≥ 2000 ㎎/㎏ Rabbit'),(CB,'LD50 > 8000 ㎎/㎏ Rabbit')])
before('- 피부부식성 또는 자극성',a11,b11,[(FE,'분진 LC50 5.05 ㎎/ℓ 4 hr Rat (OECD TG 403, GLP)'),(CB,'분진 LC50 > 4.6 ㎎/㎥ 4 hr Rat (사망없음)')])
before('- 심한 눈손상 또는 자극성',a11,b11,[(AL,'자극성 없음, Rabbit, OECD TG 404'),(FE,'토끼를 이용한 피부자극성시험(OECD TG 404) 결과 자극성 없음'),(AEA,'피부에 자극을 일으킴'),(CB,'토끼를 이용한 피부자극성시험(OECD TG 404) 결과 자극성 없음')])
before('- 호흡기과민성',a11,b11,[(AL,'자극성 없음, Rabbit, 48시간 내 완전히 가역적, OECD TG 405'),(FE,'토끼를 이용한 눈자극성시험(OECD TG 405, GLP) 결과 자극성 없음'),(AEA,'눈에 자극을 일으킴'),(CB,'토끼를 이용한 눈자극성시험(OECD TG 405) 결과 자극성 없음')])
before('- 발암성',a11,b11,[(AL,'과민성 없음, Guinea pig, 기니피그 극대화 시험(GPMT), OECD TG 406'),(FE,'기니피그를 이용한 과민성시험결과 비과민성'),(CB,'기니피그를 이용한 피부과민성시험(OECD TG 406, GLP) 결과 과민성 없음')])
before('- IARC',a11,b11,[(CB,'2')])
before('- OSHA',a11,b11,[(FE,'3'),(CB,'2B')])
before('- ACGIH',a11,b11,[(CB,'해당됨')])
before('- NTP',a11,b11,[(FE,'A4'),(CB,'A3')])
before('- 생식독성',a11,b11,[(AL,'in vitro 포유류 세포를 이용한 염색체 이상 시험: 양성(대사활성계 없음), OECD TG 473'),(FE,'복귀돌연변이시험·염색체이상시험(OECD TG 473, GLP)·in vivo comet assay 결과 음성'),(SX,'음성(복귀돌연변이시험, 염색체이상시험)'),(CB,'복귀돌연변이·유전자돌연변이·자매염색분체교환·초파리 시험 결과 음성(OECD TG 471, 476, 477, 479)')])
before('- 특정 표적장기 독성 (1회 노출)',a11,b11,[(CB,'랫드를 이용한 최기형성/발달독성 시험결과 생식독성 없음(OECD TG 414)')])
before('- 특정 표적장기 독성 (반복 노출)',a11,b11,[(AEA,'흡입시 기도를 자극함'),(FE,'호흡기계 자극을 일으킴')])
before('- 흡인유해성',a11,b11,[(FE,'반복적인 흡입노출 시 호흡기 영향(철폐증)을 줄 수 있음'),(CB,'10년 이상 반복 노출 시 만성 기관지염, 폐 기능 장애, 진폐증 등이 발생함')])
before('- 갑각류',a12,b12,[(AL,'NOEC > 50 ㎎/ℓ 96 hr Ictalurus punctatus'),(FE,'LC0 ≥ 50000 ㎎/ℓ 96 hr Danio rerio'),(SX,'LC50 100 ㎎/ℓ 96 hr'),(CB,'LC50 > 1000 ㎎/ℓ 96 hr Tribolodon hakonensis')])
before('- 조류',a12,b12,[(AL,'NOEC > 22.6 ㎎/ℓ 96 hr Acronuria sp.'),(FE,'EC50 > 100 ㎎/ℓ 48 hr'),(SX,'EC50 92.2 ㎎/ℓ 48 hr Daphnia magna'),(HAA,'LC50 24000 ㎎/ℓ 48 hr (추정치)'),(CB,'EC50 > 5600 ㎎/ℓ 24 hr Daphnia magna')])
before('나. 잔류성 및 분해성',a12,b12,[(HAA,'EC50 29.401 ㎎/ℓ 96 hr (추정치)'),(CB,'ErC50 > 10000 ㎎/ℓ 72 hr Desmodesmus subspicatus')])
before('- 분해성',a12,b12,[(SX,'log Kow 4.95 (25 ℃)'),(HAA,'log Kow -3.82 (추정값)')])
before('나. 화학물질관리법에 의한 규제',a15,b15,[(AL,'관리대상유해물질'),(None,'작업환경측정대상물질 (측정주기 : 6개월)'),(None,'특수건강진단대상물질 (진단주기 : 12개월)'),(None,'노출기준설정물질'),
                                         (FE,'관리대상유해물질'),(None,'작업환경측정대상물질 (측정주기 : 6개월)'),(None,'특수건강진단대상물질 (진단주기 : 12개월)'),(None,'노출기준설정물질'),(CB,'노출기준설정물질')])
before('바. 기타 국내 및 외국법에 의한 규제',a15,b15,[(AL,'지정폐기물'),(FE,'지정폐기물'),(SX,'지정폐기물')])
r_osha=row_of(ws,lambda t:t=='- OSHA',a11,b11); ws.cell(r_osha,2).value=None
for row,kind,items in sorted(OPS,key=lambda o:o[0],reverse=True):
    if kind=='del': delete_row(ws,row)
    else: insert_before(ws,row,items)
apply_sec2(ws)
wb.save('DY#2-WH029 EAGON WHITE 반광 MSDS.xlsx'); print('saved')
