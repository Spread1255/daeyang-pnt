import sys,glob,re,openpyxl
sys.path.insert(0,'/tmp/claude-0/ral')
from ex import txt,sections
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
# 원료명 → (CAS, 비고)
RAW=[('HC7803','27923-68-8','수지(폴리에스테르)'),('HC6011','114267-10-6','수지(폴리에스테르), CAS 파일: HC6011/M201-6'),('HC5803','114267-10-6','수지(폴리에스테르)'),
('HC6805','114267-10-6','수지(폴리에스테르)'),('PC2802','40471-09-8','수지(폴리에스테르, 외장)'),('YD013K','25036-25-3','수지(에폭시)'),
('KER3002','25036-25-3','수지(에폭시) = YD013K (사용자 확인)'),('KER3003','25036-25-3','수지(에폭시) = YD013K (사용자 확인)'),('KER3005K','25036-25-3','수지(에폭시) — CAS 파일에 없음, 금호피앤비(KER3003과 같은 제조사)·이전 배합 KD-214M 대체 → YD013K로 추정, 확인 필요'),
('AD P964','25214-38-4','CAS 파일: TC001/AD P964'),('XG603-1A','54553-90-1','경화제'),('PCA-501','6334-25-4','경화제(HAA), CAS 파일: PCA-501/XL-552'),
('LA-100','9003-49-0','레벨링제, 이전 Sheet1: LA-100/BYK-368P/L88'),('L-88','9003-49-0','레벨링제'),('BENZOIN','119-53-9','탈포제'),
('JS-250','9002-88-4','왁스(PE)'),('CERA-969','9002-88-4','왁스(PE), 이전 Sheet1에만 있음 (CERA-994와 같은 칸)'),('VISCOGEL','68953-58-2','텍스처제(유기점토)'),
('AHA-2139','9010-77-9','첨가제'),('CT-0230','60621-79-6','첨가제(아크릴 공중합체), KOSHA 결과 없음'),('SONOX626','26741-53-7','산화방지제 (CAS 파일 물질명 Irganox 242는 오류)'),
('LS-8K','21645-51-2','충전제(수산화알루미늄)'),('S#3000','1317-65-3','충전제(석회석), CAS 파일: S3000'),('SN7000','1317-65-3','충전제(석회석)'),
('KCB-02','13462-86-7','충전제(바라이트)'),('FR-767','13463-67-7','안료(이산화티타늄)'),
('HI BLACK 50L','1333-86-4','안료(카본블랙) — CAS 파일 57455-37-5는 오류 (사용자 확인)'),('RED 115','1309-37-1','안료(산화철)'),
('YELLOW 4000','5567-15-7','안료(PY83)'),('YELLOW 420','51274-00-1','안료(황색 산화철)'),('ORANGE 920P','15793-73-4','안료(PO34)'),
('BLUE 515303','147-14-8','안료(PB15)'),('VIOLET 23R','6358-30-1','안료(PV23)'),('MARINE BLUE 2504','57455-37-5','안료(PB29) = MARINE BLUE H55 (사용자 확인)'),
('REFLEXAL IV','7429-90-5','메탈릭(알루미늄)'),('PC100','7429-90-5','메탈릭(알루미늄), CAS 파일: PC 100'),('PCR 214','7429-90-5','메탈릭(알루미늄)'),
('CERA-994','9002-88-4','DFA — MSDS 제외'),('HDK V15','112945-52-5','DFA — MSDS 제외'),
('HC 6011','114267-10-6','= HC6011 (띄어쓰기 차이, S/T 조제)'),('BYK 368P','9003-49-0','= LA-100/BYK-368P/L88 (S/T 조제)'),('S 3000','1317-65-3','= S3000 (S/T 조제)'),
('RB 101A','(확인 필요)','스트럭처·함마톤 첨가제 — CAS 파일에 RB101A 이름만 있고 CAS 없음'),('101A(50%)','(확인 필요)','RB 101A 50% 희석본 — 나머지 50% 확인 필요'),
('NB 0020','(확인 필요)','CAS 파일에 없음 (S/T 조제 25%)'),('S-160(5%)','1333-86-4 (추정)','BLACK S-160(카본블랙) 5% 희석본 추정 — 나머지 95% 확인 필요')]
def kosha(cas):
    fs=glob.glob(f'/tmp/claude-0/ral/d_{cas}_*.html')
    if not fs: return None
    L,s=sections(txt(fs[0]))
    s2=s[2]; i=next(k for k,x in enumerate(s2) if x.startswith('가.')); j=next(k for k,x in enumerate(s2) if x.startswith('나.'))
    cls='; '.join(x for x in s2[i+1:j] if x!='자료없음') or '자료없음'
    h=' '.join(re.match(r'^(H\d{3})',l).group(1) for l in s2 if re.match(r'^H\d{3}',l)) or '-'
    s8=s[8]; a=next(k for k,x in enumerate(s8) if x=='국내규정'); g=next(k for k,x in enumerate(s8) if x=='ACGIH 규정'); b=next(k for k,x in enumerate(s8) if x.startswith('생물학적'))
    dom='; '.join(x.replace('TWA : ','TWA ') for x in s8[a+1:g] if x!='자료없음') or '-'
    acg='; '.join(x for x in s8[g+1:b] if x!='자료없음') or '-'
    name=s[3][2]
    return name,cls,h,dom,acg
wb=openpyxl.Workbook(); ws=wb.active; ws.title='원료 사전'
H=['원료명','CAS','KOSHA 물질명','KOSHA 유해성 분류','H코드','국내 노출기준','ACGIH','비고']
thin=Side(style='thin',color='BBBBBB'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
for j,h in enumerate(H,1):
    c=ws.cell(1,j,h); c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='E8E8E8'); c.border=B
miss=[]
for i,(n,cas,note) in enumerate(RAW,2):
    k=kosha(cas)
    if not k: miss.append(cas); k=('(KOSHA 결과 없음)','-','-','-','-')
    for j,v in enumerate([n,cas,*k,note],1):
        c=ws.cell(i,j,v); c.border=B; c.alignment=Alignment(wrap_text=True,vertical='top')
for col,w in zip('ABCDEFGH',(18,13,40,42,22,28,14,40)): ws.column_dimensions[col].width=w
ws.freeze_panes='A2'
# 제품 목록
p=wb.create_sheet('제품 목록')
P=[('DY#1-IV021','내장형 D-30871 유광','에폭시·폴리에스터','일반','완료 (임시 103)',''),
('DY#1-BL141','내장형 RAL 5005 유광','에폭시·폴리에스터','일반','완료 (번호 미정)',''),
('DY#2-WH029','외장형 EAGON WHITE 반광','폴리에스터','일반','완료 (134)',''),
('DY#4-BE005','내장형 T/X 대림 BEIGE 반광','에폭시·폴리에스터','텍스처','정보 수집 완료',''),
('DY#4-WH004','내장형 T/X HK WHITE 반광','에폭시·폴리에스터','텍스처','정보 수집 완료','RED 115 0.0035 제외'),
('DY#4-BR025','외장형 T/X PANTONE 7565U 반광','폴리에스터','텍스처','정보 수집 완료',''),
('DY#4-GY066','외장형 T/X N 2.0 반광','폴리에스터','텍스처','정보 수집 완료',''),
('DY#5-SI081','내장형 M/T 43818XL(C) 반광','에폭시·폴리에스터(추정)','메탈릭','반제품 배합 필요','M/T 43818XL 반제품 99.87 + REFLEXAL IV 0.25'),
('DY#5-SI019','내장형 HI SILVER 유광','에폭시·폴리에스터(추정)','메탈릭','반제품 배합 필요','HI SILVER 반제품 99.6 + PC100 0.4'),
('DY#5-SI060','외장형 M/T MS GRAY 반광','폴리에스터(추정)','메탈릭','BASE 배합 필요','MS GRAY BASE 99.92 + REFLEXAL IV 0.08'),
('DY#5-SI095','내장형 M/T RAL 9006 반광','에폭시·폴리에스터','메탈릭','정보 수집 완료 (알루미늄 원료 없음 — 확인 필요)','배합에 알루미늄 안료가 없음'),
('DY#5-SI094','내장형 호진 SILVER 반광','에폭시·폴리에스터','메탈릭','정보 수집 완료 (알루미늄 원료 없음 — 확인 필요)','배합에 알루미늄 안료가 없음'),
('DY#3-WH003','내장형 S/T WHITE 반광','에폭시·폴리에스터','스트럭처','조제 배합 필요','반제품(배합 있음) 99 + S/T 조제 0.8 / MARINE BLUE 2504 0.0057 제외'),
('DY#3-IV023','내장형 S/T RAL 9010 반광','에폭시·폴리에스터','스트럭처','정보 수집 완료 (KER3005K 확인 필요)','RED 115 0.008 제외'),
('DY#3-BR003','외장형 S/T EM120K 반광','폴리에스터','스트럭처','정보 수집 완료',''),
('DY#3-BL011','외장형 S/T RAL 5003 반광','폴리에스터(추정)','스트럭처','베이스·조제 배합 필요','S/T 베이스 99 + 조제 1'),
('DY#6-BE001','내장형 함마톤 BEIGE 유광','에폭시·폴리에스터','함마톤','첨가제 확인 필요','베이지 베이스(배합 있음) 99 + 101A(50%) 0.44 + S-160(5%) 0.58 / HI BLACK 0.0032 제외'),
('DY#6-BE003','내장형 함마톤 L/BEIGE 유광','에폭시·폴리에스터','함마톤','첨가제 확인 필요','베이지 베이스(BE001과 같음) 99 + 101A(50%) 0.22 + S-160(5%) 0.42'),
('DY#6-GY021','외장형 함마톤 001 GRAY 유광','폴리에스터','함마톤','첨가제 확인 필요','001 베이스(배합 있음) 99.45 + 101A(50%) 0.2 + S-160(5%) 0.5 / 작업지시서 GARY → GRAY (사용자 확인)'),
('(조제) 내장형 STRUCTURE 조제','S/T용 조제 (2013)','-','조제','RB 101A·NB 0020 확인 필요','HC 6011 40, BYK 368P 0.4, RB 101A 9, NB 0020 25, S 3000 25 (합 99.4)'),
('SI092 (파일 안 코드 DY#5-SI043)','외장형 SUPER SILVER 유광 반제품','폴리에스터(추정)','메탈릭','반제품 배합·코드 확인 필요','한빛 SILVER 반제품 96 + PCR 214 4')]
for j,h in enumerate(['코드','제품명','수지 계열 (번호 순서용)','종류','상태','비고'],1):
    c=p.cell(1,j,h); c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='E8E8E8'); c.border=B
for i,row in enumerate(P,2):
    for j,v in enumerate(row,1): c=p.cell(i,j,v); c.border=B
for col,w in zip('ABCDEF',(26,32,24,10,26,44)): p.column_dimensions[col].width=w
wb.save('MSDS 원료 사전.xlsx'); print('saved; KOSHA 없음:',miss)
