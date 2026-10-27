import copy, re
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.cell_range import CellRange
IND=' '*17
def norm(v): return re.sub(r'\s+',' ',str(v or '').strip())
def _snap(ws):
    merges=[CellRange(str(m)) for m in ws.merged_cells.ranges]
    for m in merges: ws.unmerge_cells(str(m))
    hs={r:d.height for r,d in ws.row_dimensions.items() if d.height}
    brks=[b.id for b in ws.row_breaks.brk]
    return merges,hs,brks
def _restore(ws,merges,hs,brks,at,n):
    # rows >= at shift by n (n<0: deletion of rows at..at-n-1)
    for r in list(ws.row_dimensions.keys()):
        ws.row_dimensions[r].height=None
    for r,h in hs.items():
        if n<0 and at<=r<at-n: continue
        ws.row_dimensions[r+n if r>=at else r].height=h
    for m in merges:
        if n<0:
            d0,d1=at,at-n-1                      # 지워진 행 범위
            if m.max_row<d0: pass
            elif m.min_row>d1: m.shift(row_shift=n)
            elif m.min_row>=d0 and m.max_row<=d1: continue
            else:                                # 일부만 겹치는 병합은 줄인다
                lo=min(m.min_row,d0) if m.min_row<d0 else d0
                hi=m.max_row+n if m.max_row>d1 else d0-1
                if hi<lo or (hi==lo and m.min_col==m.max_col): continue
                m=CellRange(min_col=m.min_col,min_row=lo,max_col=m.max_col,max_row=hi)
        elif m.min_row>=at: m.shift(row_shift=n)
        if m.min_row==m.max_row and m.min_col==m.max_col: continue
        ws.merge_cells(str(m))
    ws.row_breaks=RowBreak()
    for b in brks: ws.row_breaks.append(Break(id=b+n if b>=at else b))
def delete_row(ws,r):
    merges,hs,brks=_snap(ws); ws.delete_rows(r,1); _restore(ws,merges,hs,brks,r,-1)
def snap_row(ws,k):
    return {'styles':[copy.copy(ws.cell(k,c)._style) for c in range(1,6)],
            'merges':[(m.min_col,m.max_col) for m in ws.merged_cells.ranges if m.min_row==m.max_row==k],
            'height':ws.row_dimensions[k].height or 20}
def insert_before(ws,r,items,tpl=None):
    snap=tpl if isinstance(tpl,dict) else None
    tpl=None if snap else (tpl or r-1); n=len(items)
    if snap is None: snap=snap_row(ws,tpl)
    merges,hs,brks=_snap(ws)
    tmerge=snap['merges']
    ws.insert_rows(r,n); _restore(ws,merges,hs,brks,r,n)
    for i,(a,b) in enumerate(items):
        rr=r+i
        for c in range(1,6): ws.cell(rr,c)._style=copy.copy(snap['styles'][c-1])
        ws.cell(rr,1).value=(IND+a) if a else None; ws.cell(rr,2).value=b
        ws.row_dimensions[rr].height=snap['height']
        for c1,c2 in tmerge: ws.merge_cells(start_row=rr,start_column=c1,end_row=rr,end_column=c2)
def sec_range(ws,n):
    a=next(k for k in range(1,ws.max_row+1) if re.match(rf'^{n}\.\s*',norm(ws.cell(k,1).value)) and ws.cell(k,1).fill.fgColor.rgb.endswith('E6E6FA'))
    b=next((k for k in range(a+1,ws.max_row+1) if re.match(rf'^{n+1}\.\s*',norm(ws.cell(k,1).value))),ws.max_row)
    return a,b
def row_of(ws,pred,lo,hi):
    for k in range(lo,hi):
        if pred(norm(ws.cell(k,1).value)): return k
    raise KeyError(pred)
def rng(v):
    """작업지시서 백분율 → MSDS 함유량 범위 (EAGON WHITE 지정값에서 뽑은 기준, 함유량-범위-기준.md 참고)
    가운데값 c = 단위(step)로 반올림(0.5는 내림), 범위 = c-step ∽ c+step
      50 이상: 정수, ±2        (57 → 55 ∽ 59)
      1~50  : 정수, ±1        (22.5 → 21 ∽ 23, 7.56 → 7 ∽ 9, 4.8 → 4 ∽ 6, 3 → 2 ∽ 4)
      0.1~1 : 소수 첫째, ±0.1  (0.85 → 0.7 ∽ 0.9, 0.3 → 0.2 ∽ 0.4)
      0.01~0.1 : 소수 둘째, ±0.01 (0.05 → 0.04 ∽ 0.06)
      소수 셋째 자리 이상 값은 둘째 자리까지 버리고 ±0.01 (0.1392 → 0.12 ∽ 0.14)
      하한은 단위 1칸 아래로 내려가지 않음 / 0.01 미만은 KEEP()에서 제외"""
    import math
    s=repr(v); d=len(s.split('.')[1].rstrip('0')) if '.' in s else 0
    if d>=3 and v<1: v=int(v*100+1e-9)/100; st=0.01      # 1 미만 소수 셋째 자리 이상: 둘째 자리 ±0.01
    else: st=2 if v>=50 else 1 if v>=1 else 0.1 if v>=0.1 else 0.01
    u=1 if st==2 else st
    c=math.ceil(round(v/u-0.5,9))*u
    lo=max(c-st,u); hi=c+st
    dp=0 if u>=1 else (1 if u==0.1 else 2)
    return f'{lo:.{dp}f} ∽ {hi:.{dp}f}'
def KEEP(v): return v>=0.01

# 2번 유해성·위험성 (KOSHA 문구; '…' 빈칸은 채움) — IV064 틀의 16~34행 구조 기준
SEC2_CLASS=['급성 독성(경구) : 구분4','피부 부식성/피부 자극성 : 구분2','발암성 : 구분2']
SEC2_H=['H351 암을 일으킬 것으로 의심됨','H315 피부에 자극을 일으킴','H302 삼키면 유해함']
SEC2_PREV=['P201 사용 전 취급 설명서를 확보하시오.','P202 모든 안전 예방조치 문구를 읽고 이해하기 전에는 취급하지 마시오.',
 'P264 취급 후에는 손을 철저히 씻으시오.','P270 이 제품을 사용할 때에는 먹거나, 마시거나 흡연하지 마시오.','P280 보호장갑/보호의/보안경/안면보호구를(을)착용하시오.']
SEC2_RESP=['P301+P312 삼켰다면: 불편함을 느끼면 의료기관/의사의 진찰을 받으시오.','P302+P352 피부에 묻으면: 다량의 물로 씻으시오.',
 'P308+P313 노출되거나 노출이 우려되면:의학적인 조치/조언을 받으시오.','P321 적절한 처치를 하시오.','P330 입을 씻어내시오.',
 'P332+P313 피부 자극이 나타나면: 의학적인 조치/조언을 받으시오.','P362+P364 오염된 의류를 벗고 다시 사용 전 세척하시오.']
CARC_P={'P201','P202','P308+P313','P405'}  # 발암성에만 붙는 예방조치문구
def apply_sec2(ws,carc='2'):
    """carc: '2'(기본), '1A'(H350·위험), None(이산화티타늄 없음 → 발암성·H351·GHS08·발암성 전용 P문구 뺌)"""
    a,b=sec_range(ws,2)
    assert norm(ws.cell(a+1,1).value).startswith('가.') and a==15
    cls,hs=list(SEC2_CLASS),list(SEC2_H)
    prev,resp,store=list(SEC2_PREV),list(SEC2_RESP),['P405 잠금장치를 하여 저장하시오.']
    if carc=='1A':  # 발암성 1A 원료 함유(예: YELLOW P210) — 구분2 대신 1A, H350, 신호어 위험
        cls[2]='발암성 : 구분1A'; hs[0]='H350 암을 일으킬 수 있음'; ws.cell(22,2).value='위험'
    if carc is None:
        cls=cls[:2]; hs=hs[1:]
        keep=lambda t:t.split()[0] not in CARC_P
        prev=[t for t in prev if keep(t)]; resp=[t for t in resp if keep(t)]; store=[]
    def fit(label,nxt,vals):  # label 블록 행 수를 vals에 맞추고 값 채움
        r=row_of(ws,lambda t:t==label,a,b+20); e=row_of(ws,lambda t:t==nxt,r+1,b+20); n=e-r
        if len(vals)>n: insert_before(ws,e,[(None,None)]*(len(vals)-n))
        for k in range(r+n-1,r+max(len(vals),1)-1,-1): delete_row(ws,k)
        if not vals: delete_row(ws,r); return
        for k,t in enumerate(vals): ws.cell(r+k,2).value=t
    fit('저장','폐기',store); fit('대응','저장' if store else '폐기',resp); fit('예방','대응',prev)
    for k,t in enumerate(hs): ws.cell(23+k,2).value=t
    for k in range(23+len(hs),26): delete_row(ws,23+len(hs))
    for k,t in enumerate(cls): ws.cell(16+k,2).value=t
    for k in range(16+len(cls),19): delete_row(ws,16+len(cls))

# 수지(폴리에스테르·에폭시·아크릴 수지) — 1 이상이면 정수로 버린 값 ±2 (22 → 20 ∽ 24, 18.6 → 16 ∽ 20, 57 → 55 ∽ 59)
RESIN_CAS={'27923-68-8','114267-10-6','25036-25-3','40471-09-8','37953-21-2','25214-38-4'}
def rng_item(cas,v):
    if cas in RESIN_CAS and v>=1:
        c=int(v+1e-9); return f'{max(c-2,1)} ∽ {c+2}'
    return rng(v)
