import openpyxl, datetime, sys
sys.path.insert(0,'.')
from xlhelp import *
wb=openpyxl.load_workbook('/tmp/claude-0/ral/iv064.xlsx')
if '호환성 보고서' in wb.sheetnames: del wb['호환성 보고서']
ws=wb['Sheet1']
NAME='내장형 D-30871 유광'; CODE='DY#1-IV021'; NO='AA14944-0000000103'
OUT=sys.argv[1]
T={r:[ws.cell(r,c).value for c in (1,3,4)] for r in range(38,50)}
p278,p114,epx,pba,ben,pe,bar,lime,tio,pb29,y83,fe=[T[r] for r in range(38,50)]
NEW=[ # 26.6.25 작업지시서 NO 순서 그대로 (DFA 제외)
 (*p278,22),(*p114,18.6),(*epx,18.48),
 ('폴리에스테르 수지','1,3-Benzenedicarboxylic acid, polymer with 1,4-benzenedicarboxylic acid and 2,2-dimethyl-1,3-propanediol','25214-38-4',0.25),
 (*pba,0.9),(*ben,0.3),(*pe,0.2),(*lime,7.4),(*bar,11.95),(*tio,19.25),('카본블랙',None,'1333-86-4',0.0017),(*fe,0.022),(*y83,0.48)]
assert abs(sum(n[3] for n in NEW)-99.8337)<1e-9 and len(NEW)==13
NEW=[n for n in NEW if KEEP(n[3])]          # 0.01% 미만(소수점 셋째 자리 이하) 성분 제외
INC={n[0] for n in NEW}
RANGE={}   # 제품별 지정값 (수지는 rng_item에서 ±2)
if len(NEW)>12: insert_before(ws,50,[(None,None)])
for i,(a,c,d,p) in enumerate(NEW):
    r=38+i; ws.cell(r,1).value=a; ws.cell(r,3).value=c; ws.cell(r,4).value=d; ws.cell(r,5).value=RANGE.get(d) or rng_item(d,p)
    L=max(len(a),len(c or '')); ws.row_dimensions[r].height=56.25 if L>90 else 45 if L>60 else (33.75 if L>34 else 24.75)
ws['B3']=NAME; ws['B6']=f'{NAME} ({CODE})'
ws['A4'].value=ws['A4'].value.replace('AA14944-0000000102',NO)
a16,b16=sec_range(ws,16); r=row_of(ws,lambda t:t.startswith('나. 최초작성일'),a16,b16)
ws.cell(r,3).value=datetime.datetime(2024,5,9)
ws.cell(row_of(ws,lambda t:t=='개정횟수',r,ws.max_row+1),3).value='1 회'
ws.cell(row_of(ws,lambda t:t=='최종개정일자',r,ws.max_row+1),3).value=datetime.datetime(2026,1,6)
OPS=[]
a8,b8=sec_range(ws,8); a11,b11=sec_range(ws,11); a12,b12=sec_range(ws,12); a15,b15=sec_range(ws,15)
def before(label,lo,hi,items):
    items=[it for it in items if it[0] in INC or it[0] is None]
    if items: OPS.append((row_of(ws,lambda t:t==label,lo,hi),items))
CB='카본블랙'; FE='산화철'
before('ACGIH 규정',a8,b8,[(FE,'TWA -  5mg/m3'),(CB,'TWA -  3.5mg/m3   발암성 2, 흡입성')])
before('나. 적절한 공학적 관리',a8,b8,[(CB,'TWA - 3  ㎎/㎥')])
before(norm(ws.cell(row_of(ws,lambda t:t.startswith('C.I. 색소 황색 83'),a11,b11),1).value),a11,b11,[(FE,'LD50 > 10000 ㎎/㎏ Rat'),(CB,'LD50 > 8000 ㎎/㎏ Rat (OECD TG 401)')])
before('- 흡입',a11,b11,[(CB,'LD50 > 8000 ㎎/㎏ Rabbit')])
before('- 피부부식성 또는 자극성',a11,b11,[(FE,'분진 LC50 5.05 ㎎/ℓ 4 hr Rat (OECD TG 403, GLP)'),(CB,'분진 LC50 > 4.6 ㎎/㎥ 4 hr Rat (사망없음)')])
before('- 심한 눈손상 또는 자극성',a11,b11,[(FE,'토끼를 이용한 피부자극성시험(OECD TG 404) 결과 자극성 없음'),(CB,'토끼를 이용한 피부자극성시험(OECD TG 404) 결과 자극성 없음')])
before('- 호흡기과민성',a11,b11,[(FE,'토끼를 이용한 눈자극성시험(OECD TG 405, GLP) 결과 자극성 없음'),(CB,'토끼를 이용한 눈자극성시험(OECD TG 405) 결과 자극성 없음')])
before('- 발암성',a11,b11,[(FE,'기니피그를 이용한 과민성시험결과 비과민성'),(CB,'기니피그를 이용한 피부과민성시험(OECD TG 406, GLP) 결과 과민성 없음')])
before('- IARC',a11,b11,[(CB,'2')])
before('- OSHA',a11,b11,[(FE,'3'),(CB,'2B')])
before('- ACGIH',a11,b11,[(CB,'해당됨')])
before('- NTP',a11,b11,[(FE,'A4'),(CB,'A3')])
before('- 생식독성',a11,b11,[(FE,'복귀돌연변이시험·염색체이상시험(OECD TG 473, GLP)·in vivo comet assay 결과 음성'),(CB,'복귀돌연변이·유전자돌연변이·자매염색분체교환·초파리 시험 결과 음성(OECD TG 471, 476, 477, 479)')])
before('- 특정 표적장기 독성 (1회 노출)',a11,b11,[(CB,'랫드를 이용한 최기형성/발달독성 시험결과 생식독성 없음(OECD TG 414)')])
before('- 특정 표적장기 독성 (반복 노출)',a11,b11,[(FE,'호흡기계 자극을 일으킴')])
before('- 흡인유해성',a11,b11,[(FE,'반복적인 흡입노출 시 호흡기 영향(철폐증)을 줄 수 있음'),(CB,'10년 이상 반복 노출 시 만성 기관지염, 폐 기능 장애, 진폐증 등이 발생함')])
before('- 갑각류',a12,b12,[(FE,'LC0 ≥ 50000 ㎎/ℓ 96 hr Danio rerio'),(CB,'LC50 > 1000 ㎎/ℓ 96 hr Tribolodon hakonensis')])
before('- 조류',a12,b12,[(FE,'EC50 > 100 ㎎/ℓ 48 hr'),(CB,'EC50 > 5600 ㎎/ℓ 24 hr Daphnia magna')])
before('나. 잔류성 및 분해성',a12,b12,[(CB,'ErC50 > 10000 ㎎/ℓ 72 hr Desmodesmus subspicatus')])
before('나. 화학물질관리법에 의한 규제',a15,b15,[(FE,'관리대상유해물질'),(None,'작업환경측정대상물질 (측정주기 : 6개월)'),(None,'특수건강진단대상물질 (진단주기 : 12개월)'),(None,'노출기준설정물질'),(CB,'노출기준설정물질')])
before('바. 기타 국내 및 외국법에 의한 규제',a15,b15,[(FE,'지정폐기물')])
if CB in INC: r_osha=row_of(ws,lambda t:t=='- OSHA',a11,b11); ws.cell(r_osha,2).value=None
for row,items in sorted(OPS,key=lambda o:o[0],reverse=True): insert_before(ws,row,items)
apply_sec2(ws)
wb.save(OUT); print('saved',OUT)
